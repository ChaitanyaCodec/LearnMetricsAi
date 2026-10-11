"""On-demand exports for teacher-reviewed document drafts."""

from io import BytesIO
import re
from xml.sax.saxutils import escape

from django.http import HttpResponse

from docx import Document as WordDocument
from docx.shared import Inches, Pt
from pptx import Presentation
from pptx.util import Inches as PptxInches, Pt as PptxPt
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER
from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import inch
from reportlab.platypus import PageBreak, Paragraph, SimpleDocTemplate, Spacer


MAX_EXPORT_CHARACTERS = 60000
SUPPORTED_FORMATS = {"pdf", "docx", "pptx"}


def _validate_content(content):
    if not isinstance(content, str) or not content.strip():
        raise ValueError("Document content is required.")
    if len(content) > MAX_EXPORT_CHARACTERS:
        raise ValueError(f"Document content cannot exceed {MAX_EXPORT_CHARACTERS} characters.")
    return content.strip()


def _lines(content):
    return [line.rstrip() for line in content.replace("\r\n", "\n").split("\n")]


def _clean_markdown(text):
    text = re.sub(r"\*\*(.*?)\*\*", r"\1", text)
    text = re.sub(r"__(.*?)__", r"\1", text)
    text = re.sub(r"(?<!\*)\*(?!\*)(.*?)\*(?!\*)", r"\1", text)
    text = re.sub(r"`([^`]+)`", r"\1", text)
    text = re.sub(r"\[([^\]]+)\]\([^)]*\)", r"\1", text)
    return text.strip()


def _response(buffer, content_type, filename):
    response = HttpResponse(buffer.getvalue(), content_type=content_type)
    response["Content-Disposition"] = f'attachment; filename="{filename}"'
    response["Cache-Control"] = "no-store, private"
    response["X-Content-Type-Options"] = "nosniff"
    return response


def export_docx(content, filename="teacher-document.docx"):
    content = _validate_content(content)
    document = WordDocument()
    document.styles["Normal"].font.name = "Aptos"
    document.styles["Normal"].font.size = Pt(10.5)
    for line in _lines(content):
        stripped = line.strip()
        if not stripped:
            continue
        if stripped.startswith("### "):
            document.add_heading(_clean_markdown(stripped[4:]), level=3)
        elif stripped.startswith("## "):
            document.add_heading(_clean_markdown(stripped[3:]), level=2)
        elif stripped.startswith("# "):
            document.add_heading(_clean_markdown(stripped[2:]), level=1)
        elif re.match(r"^[-*+]\s+", stripped):
            document.add_paragraph(_clean_markdown(re.sub(r"^[-*+]\s+", "", stripped)), style="List Bullet")
        elif re.match(r"^\d+[.)]\s+", stripped):
            document.add_paragraph(_clean_markdown(re.sub(r"^\d+[.)]\s+", "", stripped)), style="List Number")
        else:
            document.add_paragraph(_clean_markdown(stripped))
    buffer = BytesIO()
    document.save(buffer)
    buffer.seek(0)
    return _response(buffer, "application/vnd.openxmlformats-officedocument.wordprocessingml.document", filename)


def export_pdf(content, filename="teacher-document.pdf"):
    content = _validate_content(content)
    buffer = BytesIO()
    document = SimpleDocTemplate(
        buffer, pagesize=letter, rightMargin=0.7 * inch, leftMargin=0.7 * inch,
        topMargin=0.65 * inch, bottomMargin=0.65 * inch,
        title="Teacher document draft",
    )
    base = getSampleStyleSheet()
    body = ParagraphStyle("TeacherBody", parent=base["BodyText"], leading=14, spaceAfter=7)
    h1 = ParagraphStyle("TeacherH1", parent=base["Heading1"], textColor=colors.HexColor("#1f3a5f"), spaceBefore=12, spaceAfter=8)
    h2 = ParagraphStyle("TeacherH2", parent=base["Heading2"], textColor=colors.HexColor("#28527a"), spaceBefore=10, spaceAfter=6)
    h3 = ParagraphStyle("TeacherH3", parent=base["Heading3"], spaceBefore=8, spaceAfter=5)
    story = []
    for line in _lines(content):
        stripped = line.strip()
        if not stripped:
            story.append(Spacer(1, 4))
            continue
        if stripped.startswith("### "):
            style, text = h3, stripped[4:]
        elif stripped.startswith("## "):
            style, text = h2, stripped[3:]
        elif stripped.startswith("# "):
            style, text = h1, stripped[2:]
        elif re.match(r"^[-*+]\s+", stripped):
            style, text = body, "• " + re.sub(r"^[-*+]\s+", "", stripped)
        elif re.match(r"^\d+[.)]\s+", stripped):
            style, text = body, stripped
        else:
            style, text = body, stripped
        story.append(Paragraph(escape(_clean_markdown(text)), style))
    document.build(story)
    buffer.seek(0)
    return _response(buffer, "application/pdf", filename)



def _pptx_sections(content, document_title):
    """Parse simple Markdown into topic-agnostic presentation sections."""
    sections = []
    title = "Overview"
    blocks = []
    in_code = False
    code_lines = []

    def flush():
        nonlocal blocks, title
        if blocks:
            sections.append({"title": title, "blocks": blocks})
        blocks = []

    for raw in _lines(content):
        line = raw.rstrip()
        stripped = line.strip()
        if stripped.startswith("```"):
            if in_code:
                blocks.append(("code", "\n".join(code_lines)))
                code_lines = []
                in_code = False
            else:
                in_code = True
                code_lines = []
            continue
        if in_code:
            code_lines.append(line)
            continue
        if not stripped:
            continue
        if stripped.startswith("# "):
            heading = _clean_markdown(stripped[2:])
            if not sections and not blocks and (heading.casefold() == document_title.casefold()):
                continue
            flush()
            title = heading
            continue
        if stripped.startswith("## ") or stripped.startswith("### "):
            heading = _clean_markdown(stripped.lstrip("# ").strip())
            if blocks:
                flush()
            title = heading
            continue
        if stripped.startswith("|") and stripped.endswith("|"):
            cells = [cell.strip() for cell in stripped.strip("|").split("|")]
            if all(re.fullmatch(r":?-{3,}:?", cell.replace(" ", "")) for cell in cells):
                continue
            if blocks and blocks[-1][0] == "table":
                blocks[-1][1].append(cells)
            else:
                blocks.append(("table", [cells]))
            continue
        if re.match(r"^[-*+]\s+", stripped):
            blocks.append(("bullet", _clean_markdown(re.sub(r"^[-*+]\s+", "", stripped))))
        elif re.match(r"^\d+[.)]\s+", stripped):
            blocks.append(("number", _clean_markdown(re.sub(r"^\d+[.)]\s+", "", stripped))))
        else:
            blocks.append(("text", _clean_markdown(stripped)))
    if in_code and code_lines:
        blocks.append(("code", "\n".join(code_lines)))
    flush()
    return sections or [{"title": "Overview", "blocks": [("text", _clean_markdown(content))]}]


def _pptx_add_text(slide, text, left, top, width, height, *, size=18, color=None,
                   bold=False, font="Aptos", margin=0.04):
    from pptx.dml.color import RGBColor
    from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
    shape = slide.shapes.add_textbox(left, top, width, height)
    tf = shape.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = PptxInches(margin)
    tf.margin_top = tf.margin_bottom = PptxInches(margin)
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.LEFT
    p.text = str(text)
    p.font.name = font
    p.font.size = PptxPt(size)
    p.font.bold = bold
    p.font.color.rgb = color or RGBColor(45, 55, 72)
    return shape


def _pptx_add_background(slide, presentation, index, section_title=None):
    from pptx.dml.color import RGBColor
    from pptx.enum.shapes import MSO_SHAPE
    bg = slide.background.fill
    bg.solid()
    bg.fore_color.rgb = RGBColor(248, 250, 253)
    # Slim brand rail and footer are intentionally theme-wide, not topic-specific.
    rail = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, 0, 0, PptxInches(0.16), presentation.slide_height
    )
    rail.fill.solid()
    rail.fill.fore_color.rgb = RGBColor(36, 83, 145)
    rail.line.fill.background()
    footer = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, PptxInches(0.42), PptxInches(7.13),
        PptxInches(12.45), PptxInches(0.012)
    )
    footer.fill.solid()
    footer.fill.fore_color.rgb = RGBColor(220, 228, 239)
    footer.line.fill.background()
    _pptx_add_text(
        slide, f"LEARNMETRICS AI  •  {index:02d}", PptxInches(0.48),
        PptxInches(7.16), PptxInches(4.5), PptxInches(0.22),
        size=8, color=RGBColor(105, 119, 139), bold=True
    )


def _pptx_add_visual(slide, presentation, section_title, blocks):
    """Create a small editable visual from the section's own terms."""
    from pptx.dml.color import RGBColor
    from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
    from pptx.enum.text import PP_ALIGN

    combined = (section_title + " " + " ".join(str(b[1]) for b in blocks)).casefold()
    left, top, width = PptxInches(8.75), PptxInches(1.75), PptxInches(3.65)
    panel = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, PptxInches(4.55))
    panel.fill.solid()
    panel.fill.fore_color.rgb = RGBColor(235, 242, 251)
    panel.line.color.rgb = RGBColor(219, 229, 243)

    if any(word in combined for word in ("versus", " vs ", "comparison", "difference", "compare")):
        labels = ["Option A", "Option B"]
        candidates = []
        for kind, value in blocks:
            if kind in ("bullet", "number", "text") and value and len(candidates) < 2:
                candidates.append(value[:24])
        labels = (candidates + labels)[:2]
        for i, label in enumerate(labels):
            x = PptxInches(9.05 + i * 1.65)
            box = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, PptxInches(2.75), PptxInches(1.42), PptxInches(1.12))
            box.fill.solid()
            box.fill.fore_color.rgb = RGBColor(36, 83, 145) if i == 0 else RGBColor(18, 135, 126)
            box.line.fill.background()
            tf = box.text_frame
            tf.clear()
            tf.word_wrap = True
            tf.paragraphs[0].text = label
            tf.paragraphs[0].alignment = PP_ALIGN.CENTER
            tf.paragraphs[0].font.name = "Aptos"
            tf.paragraphs[0].font.size = PptxPt(12)
            tf.paragraphs[0].font.bold = True
            tf.paragraphs[0].font.color.rgb = RGBColor(255, 255, 255)
        _pptx_add_text(slide, "COMPARE", PptxInches(9.1), PptxInches(2.0), PptxInches(3.0), PptxInches(0.4), size=12, color=RGBColor(36, 83, 145), bold=True)
    elif any(word in combined for word in ("process", "workflow", "steps", "procedure", "pipeline", "execution", "algorithm")):
        labels = []
        for kind, value in blocks:
            if kind in ("bullet", "number", "text") and value:
                labels.append(value[:18])
            if len(labels) == 3:
                break
        labels = labels or ["Input", "Process", "Output"]
        for i, label in enumerate(labels):
            y = PptxInches(2.15 + i * 1.05)
            box = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, PptxInches(9.15), y, PptxInches(2.7), PptxInches(0.62))
            box.fill.solid()
            box.fill.fore_color.rgb = [RGBColor(36,83,145), RGBColor(18,135,126), RGBColor(94,103,190)][i % 3]
            box.line.fill.background()
            tf = box.text_frame
            tf.clear()
            tf.word_wrap = True
            tf.paragraphs[0].text = label
            tf.paragraphs[0].alignment = PP_ALIGN.CENTER
            tf.paragraphs[0].font.name = "Aptos"
            tf.paragraphs[0].font.size = PptxPt(11)
            tf.paragraphs[0].font.color.rgb = RGBColor(255,255,255)
            if i < len(labels)-1:
                _pptx_add_text(slide, "↓", PptxInches(10.1), y + PptxInches(0.62), PptxInches(0.5), PptxInches(0.35), size=18, color=RGBColor(36,83,145), bold=True)
    elif any(word in combined for word in ("architecture", "component", "structure", "system", "ecosystem", "framework")):
        labels = []
        for kind, value in blocks:
            if kind in ("bullet", "number") and value:
                labels.append(value[:18])
            if len(labels) >= 4:
                break
        labels = labels or ["Core concept", "Supporting part", "Output"]
        coords = [(9.65,2.15),(8.98,3.35),(10.35,3.35),(9.65,4.55)]
        center = slide.shapes.add_shape(MSO_SHAPE.OVAL, PptxInches(9.78), PptxInches(3.05), PptxInches(1.3), PptxInches(0.72))
        center.fill.solid(); center.fill.fore_color.rgb = RGBColor(36,83,145); center.line.fill.background()
        center.text_frame.text = "Core"
        center.text_frame.paragraphs[0].alignment = PP_ALIGN.CENTER
        center.text_frame.paragraphs[0].font.size = PptxPt(12)
        center.text_frame.paragraphs[0].font.bold = True
        center.text_frame.paragraphs[0].font.color.rgb = RGBColor(255,255,255)
        for i, label in enumerate(labels[:4]):
            x,y=coords[i]
            node=slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,PptxInches(x),PptxInches(y),PptxInches(1.55),PptxInches(0.48))
            node.fill.solid(); node.fill.fore_color.rgb = RGBColor(18,135,126) if i%2 else RGBColor(94,103,190); node.line.fill.background()
            node.text_frame.text=label
            node.text_frame.paragraphs[0].alignment=PP_ALIGN.CENTER
            node.text_frame.paragraphs[0].font.size=PptxPt(8)
            node.text_frame.paragraphs[0].font.color.rgb=RGBColor(255,255,255)
    else:
        # Topic-agnostic editorial illustration built from editable shapes.
        colorset = [RGBColor(36,83,145), RGBColor(18,135,126), RGBColor(94,103,190)]
        for i, (x, y, d) in enumerate([(9.25,2.35,1.12),(10.65,2.95,0.78),(9.7,4.05,0.95)]):
            shape = slide.shapes.add_shape(MSO_SHAPE.OVAL, PptxInches(x), PptxInches(y), PptxInches(d), PptxInches(d))
            shape.fill.solid(); shape.fill.fore_color.rgb = colorset[i]; shape.line.fill.background()
        _pptx_add_text(slide, "KEY IDEA", PptxInches(9.05), PptxInches(1.95), PptxInches(3.0), PptxInches(0.35), size=12, color=RGBColor(36,83,145), bold=True)


def export_pptx(content, filename="teacher-document.pptx", title=None):
    """Create a generic, theme-based presentation from edited Markdown content."""
    from pptx.dml.color import RGBColor
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

    content = _validate_content(content)
    document_title = _clean_markdown((title or "").strip()) or "Teacher Presentation"
    presentation = Presentation()
    presentation.slide_width = PptxInches(13.333)
    presentation.slide_height = PptxInches(7.5)
    presentation.core_properties.title = document_title
    presentation.core_properties.subject = "Teacher-created educational presentation"
    presentation.core_properties.author = "LearnMetrics AI"

    # Cover slide: clean editorial design with an editable geometric illustration.
    cover = presentation.slides.add_slide(presentation.slide_layouts[6])
    cover.background.fill.solid()
    cover.background.fill.fore_color.rgb = RGBColor(20, 39, 72)
    accent = cover.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, PptxInches(0.22), presentation.slide_height)
    accent.fill.solid(); accent.fill.fore_color.rgb = RGBColor(29, 183, 166); accent.line.fill.background()
    circle1 = cover.shapes.add_shape(MSO_SHAPE.OVAL, PptxInches(9.1), PptxInches(1.15), PptxInches(2.7), PptxInches(2.7))
    circle1.fill.solid(); circle1.fill.fore_color.rgb = RGBColor(29, 183, 166); circle1.fill.transparency = 12; circle1.line.fill.background()
    circle2 = cover.shapes.add_shape(MSO_SHAPE.OVAL, PptxInches(10.45), PptxInches(3.25), PptxInches(1.65), PptxInches(1.65))
    circle2.fill.solid(); circle2.fill.fore_color.rgb = RGBColor(100, 116, 210); circle2.fill.transparency = 8; circle2.line.fill.background()
    _pptx_add_text(cover, "LEARNMETRICS AI  /  TEACHING STUDIO", PptxInches(0.9), PptxInches(1.05), PptxInches(7.4), PptxInches(0.35), size=11, color=RGBColor(92, 224, 209), bold=True)
    _pptx_add_text(cover, document_title, PptxInches(0.88), PptxInches(2.0), PptxInches(7.6), PptxInches(2.05), size=31, color=RGBColor(255,255,255), bold=True)
    _pptx_add_text(cover, "A structured learning presentation", PptxInches(0.92), PptxInches(4.55), PptxInches(6.6), PptxInches(0.5), size=16, color=RGBColor(205,218,236))
    _pptx_add_text(cover, "Prepared for classroom teaching  •  Review before delivery", PptxInches(0.92), PptxInches(6.45), PptxInches(8.4), PptxInches(0.3), size=10, color=RGBColor(205,218,236))

    sections = _pptx_sections(content, document_title)
    slide_number = 1
    for section in sections:
        blocks = section["blocks"]
        # Render each Markdown table on its own slide for readable sizing.
        normal_blocks = []
        tables = []
        for block in blocks:
            (tables if block[0] == "table" else normal_blocks).append(block)
        content_groups = []
        pending = []
        for block in normal_blocks:
            if block[0] == "code":
                if pending:
                    content_groups.append(pending)
                    pending = []
                content_groups.append([block])
            else:
                pending.append(block)
                if len(pending) >= 4:
                    content_groups.append(pending)
                    pending = []
        if pending:
            content_groups.append(pending)
        if not content_groups and not tables:
            content_groups = [[]]
        for group_index, group in enumerate(content_groups):
            slide = presentation.slides.add_slide(presentation.slide_layouts[6])
            slide_number += 1
            _pptx_add_background(slide, presentation, slide_number, section["title"])
            _pptx_add_text(slide, section["title"] + (" (continued)" if group_index else ""),
                           PptxInches(0.55), PptxInches(0.48), PptxInches(11.9), PptxInches(0.72),
                           size=25, color=RGBColor(28,48,78), bold=True)
            # Accent line under the heading.
            line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, PptxInches(0.58), PptxInches(1.28), PptxInches(0.9), PptxInches(0.055))
            line.fill.solid(); line.fill.fore_color.rgb = RGBColor(29,183,166); line.line.fill.background()

            y = 1.58
            for kind, value in group:
                if kind == "code":
                    card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, PptxInches(0.62), PptxInches(y), PptxInches(7.65), PptxInches(2.5))
                    card.fill.solid(); card.fill.fore_color.rgb = RGBColor(25,35,52); card.line.fill.background()
                    tf = card.text_frame; tf.clear(); tf.word_wrap = True
                    tf.margin_left = tf.margin_right = PptxInches(0.18)
                    tf.margin_top = tf.margin_bottom = PptxInches(0.12)
                    for i, code_line in enumerate(value.splitlines()[:14]):
                        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
                        p.text = code_line[:105]
                        p.font.name = "Consolas"; p.font.size = PptxPt(10)
                        p.font.color.rgb = RGBColor(230,239,249)
                    y += 2.62
                else:
                    if kind in ("bullet", "number"):
                        marker = "•" if kind == "bullet" else "›"
                        text = f"{marker}  {value}"
                    else:
                        text = value
                    estimated = max(1, len(text) // 68 + 1)
                    height = min(0.82, max(0.42, estimated * 0.24 + 0.15))
                    if y + height > 6.45:
                        break
                    _pptx_add_text(slide, text, PptxInches(0.72), PptxInches(y), PptxInches(7.55), PptxInches(height),
                                   size=15 if len(text) < 100 else 13,
                                   color=RGBColor(55,68,88), bold=(kind == "number"))
                    y += height + 0.08
            _pptx_add_visual(slide, presentation, section["title"], group)
        for table_index, (_, rows) in enumerate(tables):
            if not rows:
                continue
            slide = presentation.slides.add_slide(presentation.slide_layouts[6])
            slide_number += 1
            _pptx_add_background(slide, presentation, slide_number, section["title"])
            _pptx_add_text(slide, section["title"] + " — Table", PptxInches(0.55), PptxInches(0.48),
                           PptxInches(11.9), PptxInches(0.72), size=24, color=RGBColor(28,48,78), bold=True)
            table_rows = rows[:9]
            max_cols = min(max(len(r) for r in table_rows), 5)
            table = slide.shapes.add_table(len(table_rows), max_cols, PptxInches(0.65), PptxInches(1.65),
                                           PptxInches(12.0), PptxInches(min(4.8, 0.48 * len(table_rows)))).table
            for ri, row in enumerate(table_rows):
                for ci in range(max_cols):
                    cell = table.cell(ri, ci)
                    cell.text = row[ci][:180] if ci < len(row) else ""
                    cell.margin_left = cell.margin_right = PptxInches(0.08)
                    cell.margin_top = cell.margin_bottom = PptxInches(0.04)
                    cell.fill.solid()
                    cell.fill.fore_color.rgb = RGBColor(36,83,145) if ri == 0 else (RGBColor(239,244,251) if ri % 2 else RGBColor(255,255,255))
                    for p in cell.text_frame.paragraphs:
                        p.font.name = "Aptos"
                        p.font.size = PptxPt(11 if len(table_rows) < 7 else 9)
                        p.font.bold = ri == 0
                        p.font.color.rgb = RGBColor(255,255,255) if ri == 0 else RGBColor(45,55,72)

    buffer = BytesIO()
    presentation.save(buffer)
    buffer.seek(0)
    return _response(buffer, "application/vnd.openxmlformats-officedocument.presentationml.presentation", filename)

def export_document(content, file_format, filename=None, title=None):
    if file_format not in SUPPORTED_FORMATS:
        raise ValueError("Unsupported export format. Choose PDF, DOCX, or PPTX.")
    exporters = {"pdf": export_pdf, "docx": export_docx, "pptx": export_pptx}
    default_name = f"teacher-document.{file_format}"
    safe_filename = filename or default_name
    if not safe_filename.lower().endswith(f".{file_format}"):
        safe_filename = f"{safe_filename}.{file_format}"
    if file_format == "pptx":
        return export_pptx(content, filename=safe_filename, title=title)
    return exporters[file_format](content, filename=safe_filename)
