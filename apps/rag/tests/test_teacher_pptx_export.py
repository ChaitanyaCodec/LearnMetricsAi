from io import BytesIO

from django.test import SimpleTestCase
from pptx import Presentation

from apps.rag.teacher.exporters import export_document, export_pptx


class GenericTeacherPptxExportTests(SimpleTestCase):
    def _presentation(self, response):
        self.assertEqual(
            response["Content-Type"],
            "application/vnd.openxmlformats-officedocument.presentationml.presentation",
        )
        self.assertEqual(response["Cache-Control"], "no-store, private")
        return Presentation(BytesIO(response.content))

    def _all_text(self, presentation):
        values = []
        for slide in presentation.slides:
            for shape in slide.shapes:
                if getattr(shape, "has_text_frame", False):
                    values.append(shape.text)
                if getattr(shape, "has_table", False):
                    for row in shape.table.rows:
                        values.extend(cell.text for cell in row.cells)
        return "\n".join(values)

    def test_exports_non_spark_topic_with_supplied_title_and_theme(self):
        response = export_pptx(
            "# Photosynthesis\n\n## Process\n- Light is absorbed by chlorophyll.\n"
            "- Plants convert light energy into chemical energy.\n"
            "- Oxygen is released as a by-product.\n\n"
            "## Key terms\n- Chlorophyll\n- Glucose\n- Carbon dioxide",
            title="Photosynthesis: How Plants Make Food",
        )
        presentation = self._presentation(response)
        self.assertGreaterEqual(len(presentation.slides), 3)
        self.assertEqual(presentation.slide_width, 12191695)  # 13.333 inches, approx.
        text = self._all_text(presentation)
        self.assertIn("Photosynthesis: How Plants Make Food", text)
        self.assertIn("Light is absorbed by chlorophyll.", text)
        self.assertIn("Chlorophyll", text)
        self.assertNotIn("Spark", text)

    def test_export_document_uses_current_title_for_pptx_cover(self):
        response = export_document(
            "## Learning objectives\n- Explain normalization.\n- Identify normal forms.",
            "pptx",
            filename="database-normalization.pptx",
            title="Database Normalization",
        )
        presentation = self._presentation(response)
        text = self._all_text(presentation)
        self.assertIn("Database Normalization", text)
        self.assertIn("Explain normalization.", text)

    def test_markdown_comparison_table_is_included(self):
        response = export_pptx(
            "## Compare approaches\n\n| Feature | Approach A | Approach B |\n"
            "|---|---|---|\n| Cost | Low | High |\n| Speed | Medium | Fast |",
            title="Comparison of Approaches",
        )
        presentation = self._presentation(response)
        text = self._all_text(presentation)
        self.assertIn("Feature", text)
        self.assertIn("Approach A", text)
        self.assertIn("Medium", text)
